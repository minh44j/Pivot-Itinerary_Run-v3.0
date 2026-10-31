"""Generate the "Manual itinerary" workbook (xlsx) that Google Sheets converts on upload.

The layout comes from the constants in manual_sheet.py — the same ones the runner reads — so the sheet and
the runner cannot drift apart. Run:   python tools/build_manual_sheet.py out.xlsx
then upload it to Drive (converted to a Google Sheet), share it with the automation's mailbox as Editor, and
put its file id in the GitHub secret MANUAL_SHEET_ID. See manual_form/README.md.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import manual_sheet as M

INK, SLATE, GOLD = "1D1D1F", "5B6470", "C9A84C"
FONT = "Arial"
thin = Side(style="thin", color="D9D4C3")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT = PatternFill("solid", fgColor="FFFDF5")
DARK = PatternFill("solid", fgColor="1E1E20")
GREY = PatternFill("solid", fgColor="F5F5F7")
GOLDF = PatternFill("solid", fgColor=GOLD)


def f(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("color", INK)
    return Font(**kw)


def section(ws, row, text, last_col="L"):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    c = ws[f"A{row}"]
    c.value = text
    c.font = f(bold=True, size=12, color="FFFFFF")
    c.fill = DARK
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = 24


def label(ws, ref, text, required=False):
    c = ws[ref]
    c.value = text + ("  *" if required else "")
    c.font = f(bold=required, size=11)
    c.alignment = Alignment(vertical="center", indent=1)


def inp(ws, ref, fmt="@", align="left"):
    c = ws[ref]
    c.fill, c.border = INPUT, BOX
    c.number_format = fmt
    c.font = f(size=11)
    c.alignment = Alignment(horizontal=align, vertical="center")
    return c


def head(ws, ref, text):
    c = ws[ref]
    c.value = text
    c.font = f(bold=True, size=10, color=INK)
    c.fill = GREY
    c.border = BOX
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = M.SHEET_TAB
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GOLD

    ws.merge_cells("A1:L1")
    ws["A1"] = "Pivot — Manual itinerary"
    ws["A1"].font = f(bold=True, size=20)
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:L2")
    ws["A2"] = ("Type the ticket exactly as printed. Fill the cream cells, choose BUILD NOW at the bottom, "
                "and the PDF arrives in the e-mail box you typed in B9 within a few minutes. "
                "Nothing goes to a client.   *  = required")
    ws["A2"].font = f(size=10, color=SLATE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30

    # 1 — the booking
    section(ws, 4, "1   The booking")
    rows = (("PNR", "PNR", True, "@"), ("Agency reference", "Agency reference", False, "@"),
            ("CRS reference", "CRS reference", False, "@"), ("Booked on", "Booked on  (empty if the ticket has none)", False, "dd mmm yyyy"))
    for key, text, req, fmt in rows:
        r, c = M.BOOKING_CELLS[key]
        label(ws, f"A{r}", text, req)
        inp(ws, f"{c}{r}", fmt)
        ws.row_dimensions[r].height = 24
    r, c = M.EMAIL_CELL
    label(ws, f"A{r}", "Send the PDF to (your e-mail)", True)
    inp(ws, f"{c}{r}")
    ws[f"{c}{r}"] = "cs@pivottmc.com"   # the executives share the cs@ mailbox; change only to send elsewhere
    ws.row_dimensions[r].height = 24
    ws.merge_cells(f"B{r}:D{r}")
    for col in "CD":
        ws[f"{col}{r}"].border = BOX

    # 2 — passengers
    section(ws, 11, "2   Passengers  (up to 4 — Passenger 1 is required)")
    hr = M.PAX_FIRST_ROW - 1
    ws[f"A{hr}"].fill = GREY
    for field, col in M.PAX_COLS.items():
        head(ws, f"{col}{hr}", {"Title": "Title", "Given name(s)": "Given name(s)\n(as on the ticket)", "Surname": "Surname\n(as on the ticket)",
                                "Ticket number": "Ticket number\n(digits only)"}[field])
    ws.row_dimensions[hr].height = 34
    for n in range(1, M.MAX_PASSENGERS + 1):
        r = M.PAX_FIRST_ROW + n - 1
        label(ws, f"A{r}", f"Passenger {n}", n == 1)
        for field, col in M.PAX_COLS.items():
            inp(ws, f"{col}{r}", "@")
        ws.row_dimensions[r].height = 24

    # 3 — flights
    section(ws, 18, "3   Flights  (in travel order, up to 6 — a connection is a separate flight)", last_col="V")
    hr = M.FLIGHT_FIRST_ROW - 1
    ws[f"A{hr}"].fill = GREY
    hints = {"Flight number": "Flight number\n(SV 774)", "Operated by (airline)": "Airline\n(Saudia)", "From (airport code)": "From\ncode (RUH)",
             "From city": "From city\n(Riyadh)", "To (airport code)": "To\ncode (COK)", "To city": "To city\n(Cochin)",
             "Departure date": "Departure\ndate", "Departure time": "Departure\ntime (18:05)", "Arrival date": "Arrival\ndate",
             "Arrival time": "Arrival\ntime (22:40)", "Cabin class": "Cabin\nclass", "Departure terminal": "Dep.\nterminal",
             "Arrival terminal": "Arr.\nterminal", "Duration": "Duration\n(4h 45m)", "Cabin baggage": "Cabin bag\n(7kg)",
             "Checked baggage": "Checked bag\n(2 pieces / 30kg)", "Airline reference": "Airline ref.\n(only if different\nfrom the PNR)"}
    for field, col in M.FLIGHT_COLS.items():
        head(ws, f"{col}{hr}", hints[field])
    for k, col in enumerate(M.SEAT_COLS, 1):
        head(ws, f"{col}{hr}", f"Seat\npassenger {k}\n(14A)")
    ws.row_dimensions[hr].height = 50
    fmts = {"Departure date": "dd mmm yyyy", "Arrival date": "dd mmm yyyy", "Departure time": "hh:mm", "Arrival time": "hh:mm"}
    for n in range(1, M.MAX_FLIGHTS + 1):
        r = M.FLIGHT_FIRST_ROW + n - 1
        label(ws, f"A{r}", f"Flight {n}", n == 1)
        for field, col in M.FLIGHT_COLS.items():
            inp(ws, f"{col}{r}", fmts.get(field, "@"), "center" if field in fmts or "code" in field else "left")
        for col in M.SEAT_COLS:
            inp(ws, f"{col}{r}", "@", "center")
        ws.row_dimensions[r].height = 22

    # 4 — build it
    section(ws, 27, "4   Build it")
    r, c = M.TRIGGER_CELL
    label(ws, f"A{r}", "When everything is typed, choose  →", True)
    t = inp(ws, f"{c}{r}", "@", "center")
    t.fill, t.font = GOLDF, f(bold=True, size=12)
    ws.row_dimensions[r].height = 28
    r, c = M.STATUS_CELL
    label(ws, f"A{r}", "What happened")
    ws.merge_cells(f"B{r}:L{r}")
    s = ws[f"{c}{r}"]
    s.value = "Nothing built yet."
    s.font = f(size=11, color=SLATE)
    s.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    s.border = BOX
    ws.row_dimensions[r].height = 92

    # widths
    ws.column_dimensions["A"].width = 38
    widths = {"B": 18, "C": 22, "D": 16, "E": 18, "F": 14, "G": 16, "H": 14, "I": 13, "J": 14, "K": 13, "L": 17,
              "M": 11, "N": 11, "O": 12, "P": 12, "Q": 17, "R": 16, "S": 12, "T": 12, "U": 12, "V": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    # validations (dropdowns + dates; the runner re-checks everything and e-mails a plain fix list)
    def dv(rng, **kw):
        v = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
        ws.add_data_validation(v)
        v.add(rng)

    pf, pl = M.PAX_FIRST_ROW, M.PAX_FIRST_ROW + M.MAX_PASSENGERS - 1
    ff, fl = M.FLIGHT_FIRST_ROW, M.FLIGHT_FIRST_ROW + M.MAX_FLIGHTS - 1
    dv(f"B{pf}:B{pl}", type="list", formula1='"' + ",".join(M.TITLES) + '"', errorTitle="Title", error="Pick a title from the list.")
    dv(f"L{ff}:L{fl}", type="list", formula1='"' + ",".join(M.CABINS) + '"', errorTitle="Cabin class", error="Pick a cabin class from the list.")
    dv(f"{M.TRIGGER_CELL[1]}{M.TRIGGER_CELL[0]}", type="list", formula1='"' + M.TRIGGER_TEXT + '"', errorTitle="Build", error="Choose BUILD NOW from the list.")
    for rng in (f"B{M.BOOKING_CELLS['Booked on'][0]}", f"H{ff}:H{fl}", f"J{ff}:J{fl}"):
        dv(rng, type="date", operator="greaterThan", formula1="36526", errorTitle="Date",
           error="Type a real date, for example 17 Dec 2026 (or pick it from the calendar).")

    # Help
    hp = wb.create_sheet(M.HELP_TAB)
    hp.sheet_view.showGridLines = False
    hp.column_dimensions["A"].width = 120
    lines = [
        ("How to use this sheet", True),
        ("1.  Go to the “New booking” tab and fill the cream cells from the ticket. Only PNR, e-mail, Passenger 1 and Flight 1 are required.", False),
        ("2.  Passengers: title, given names and surname in separate boxes, and the ticket number as digits. Up to 4.", False),
        ("3.  Flights: one row per flight, in travel order. A connection is a separate flight. Times are LOCAL at each airport, 24-hour (18:05). Up to 6.", False),
        ("4.  Seats: one box per passenger, in the order you typed them. Leave empty if no seat is assigned.", False),
        ("5.  Choose BUILD NOW (cell B28). Within a few minutes the PDF arrives in your inbox and “What happened” says DONE, and the sheet clears itself for the next booking.", False),
        ("6.  If something is wrong, nothing is built: “What happened” and an e-mail list exactly what to fix. Fix it in place and choose BUILD NOW again.", False),
        ("7.  Always check the PDF against the airline ticket before you forward it to the client.", False),
        ("", False),
        ("Rules", True),
        ("•  Type what is printed. Nothing is guessed: a date or time it cannot read is refused, and an empty “Booked on” prints N/A.", False),
        ("•  PNR 5–8 letters/digits · ticket 10–14 digits · airport codes 3 letters · seats like 14A · arrival not before departure.", False),
        ("•  Two airlines on one ticket? Type the second airline’s own reference in “Airline ref.” on that flight; both references then print.", False),
        ("•  One person at a time: the sheet holds a single booking. Wait for DONE (it clears) before starting the next.", False),
        ("•  Do not rename tabs or move cells — the automation reads fixed positions.", False),
        ("", False),
        ("What “What happened” can say", True),
        ("•  Building…   —   it picked your request up (takes about a minute).", False),
        ("•  DONE …   —   the PDF was e-mailed to the address in B9 (copy to cs@).", False),
        ("•  NOT BUILT …   —   a list of things to fix follows; your typing is still here.", False),
        ("•  ERROR …   —   something unexpected; tell the admin. Nothing was sent.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = hp[f"A{i}"]
        c.value = text
        c.font = f(bold=bold, size=14 if bold else 11, color=INK if bold else "333333")
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # History
    hs = wb.create_sheet(M.HISTORY_TAB)
    for i, t in enumerate(("When", "PNR", "Result", "Sent to"), 1):
        c = hs.cell(row=1, column=i, value=t)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = DARK
    for col, w in zip("ABCD", (20, 14, 60, 32)):
        hs.column_dimensions[col].width = w
    hs.freeze_panes = "A2"
    wb.save(path)
    return path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "manual_itinerary.xlsx"
    print("wrote", build(out))
